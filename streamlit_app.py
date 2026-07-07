import io
import os
import copy
import datetime
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

# ============================================================
#  GIOVEN TRAVEL · Excel Toolkit
#  Design giovane, arioso, mobile-first
# ============================================================

st.set_page_config(
    page_title="GIOVEN TRAVEL",
    page_icon="✈️",
    layout="centered",
    initial_sidebar_state="collapsed",
)

# ------------------------------------------------------------
#  STILE
# ------------------------------------------------------------
st.markdown(
    """
    <style>
    html, body, [class*="css"] {
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "Nunito",
                     Roboto, "PingFang SC", "Microsoft YaHei", sans-serif;
    }
    [data-testid="stHeader"] { background: transparent; }
    #MainMenu, footer, [data-testid="stToolbar"] { visibility: hidden; }

    .block-container { max-width: 720px; padding-top: 1rem; padding-bottom: 4rem; }

    /* ---------- HERO ---------- */
    .hero {
        position: relative;
        background: linear-gradient(135deg,#a78bfa 0%,#8b5cf6 40%,#ec4899 100%);
        border-radius: 26px;
        padding: 26px 26px 24px;
        color:#fff;
        box-shadow: 0 18px 40px rgba(139,92,246,.32);
        margin-bottom: 26px;
        overflow: hidden;
    }
    .hero::after {
        content:""; position:absolute; right:-40px; top:-40px;
        width:160px; height:160px; border-radius:50%;
        background: rgba(255,255,255,.15);
    }
    .hero h1 { margin:0; font-size:1.9rem; font-weight:800; letter-spacing:.3px; }
    .hero p  { margin:6px 0 0; font-size:.95rem; opacity:.95; max-width:90%; }

    /* ---------- SEZIONI (leggere, senza scatole pesanti) ---------- */
    .sec { display:flex; align-items:center; gap:9px; margin: 26px 0 12px; }
    .sec .n {
        width:26px; height:26px; border-radius:9px;
        background:#ede9fe; color:#7c3aed; font-weight:800; font-size:.85rem;
        display:flex; align-items:center; justify-content:center;
    }
    .sec .t { font-size:1.05rem; font-weight:750; color:#1f2540; }
    .sec .t small { display:block; font-size:.72rem; color:#94a3b8; font-weight:500; }

    /* ---------- SELETTORE A CARTE (pulsanti modalità) ---------- */
    div[data-testid="column"] .stButton button {
        min-height: 88px; border-radius: 18px; font-weight: 700;
        font-size: 1rem; line-height:1.35; white-space: pre-line;
        transition: all .15s ease; padding: 10px;
    }
    /* inattivo */
    .stButton button[kind="secondary"] {
        background:#faf9ff; border:1.6px solid #ece7ff; color:#5b21b6;
    }
    .stButton button[kind="secondary"]:hover {
        border-color:#c4b5fd; background:#f5f2ff; transform: translateY(-2px);
    }
    /* attivo */
    .stButton button[kind="primary"] {
        background: linear-gradient(135deg,#8b5cf6,#7c3aed); color:#fff;
        border:none; box-shadow: 0 10px 24px rgba(124,58,237,.34);
        transform: translateY(-2px);
    }

    /* ---------- UPLOADER ---------- */
    [data-testid="stFileUploader"] section {
        border:2px dashed #c9bcff; border-radius:16px; background:#faf9ff; padding:14px;
    }
    [data-testid="stFileUploader"] section:hover { border-color:#8b5cf6; }

    /* ---------- STAT RISULTATO ---------- */
    .stat {
        display:flex; align-items:center; gap:14px;
        background: linear-gradient(135deg,#f5f3ff,#fce7f3);
        border:1px solid #eaddff; border-radius:18px; padding:16px 20px; margin-top:14px;
    }
    .stat .big { font-size:2.1rem; font-weight:800;
        background:linear-gradient(135deg,#7c3aed,#ec4899);
        -webkit-background-clip:text; -webkit-text-fill-color:transparent; line-height:1; }
    .stat .lbl { font-size:.9rem; color:#6b21a8; }
    .stat .lbl b { color:#4c1d95; }

    /* ---------- DOWNLOAD ---------- */
    .stDownloadButton button {
        background: linear-gradient(135deg,#7c3aed,#ec4899) !important;
        color:#fff !important; border:none !important; width:100%;
        border-radius:16px !important; font-weight:800 !important;
        padding:.8rem 1rem !important; font-size:1.02rem !important;
        box-shadow: 0 12px 26px rgba(124,58,237,.32) !important;
    }
    .stDownloadButton button:hover { filter:brightness(1.06); transform: translateY(-1px); }

    /* ---------- EXPANDER (opzioni) ---------- */
    [data-testid="stExpander"] { border:none !important; }
    [data-testid="stExpander"] summary { border-radius:12px; background:#f6f4ff; font-weight:600; }

    /* chip multiselect: angoli morbidi */
    [data-baseweb="tag"] { border-radius:9px !important; }

    /* ---------- RESPONSIVE ---------- */
    @media (max-width: 640px) {
        .block-container { padding-left:.8rem; padding-right:.8rem; }
        .hero { padding:22px 20px; border-radius:22px; }
        .hero h1 { font-size:1.55rem; }
        .hero p  { font-size:.86rem; max-width:100%; }
        div[data-testid="column"] .stButton button { min-height: 76px; font-size:.92rem; }
        .stat .big { font-size:1.8rem; }
    }
    /* ---------- DARK ---------- */
    @media (prefers-color-scheme: dark) {
        .sec .t { color:#f1f5f9; }
        .stat { background:linear-gradient(135deg,#2a2140,#3a1d33); border-color:#4c1d95; }
        .stat .lbl { color:#e9d5ff; } .stat .lbl b{ color:#fff; }
        .stButton button[kind="secondary"]{ background:#1e1b2e; border-color:#3b2f5e; color:#ddd6fe; }
        [data-testid="stFileUploader"] section{ background:#1a1730; border-color:#5b21b6; }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ------------------------------------------------------------
#  HERO
# ------------------------------------------------------------
st.markdown(
    """
    <div class="hero">
        <h1>✈️ GIOVEN TRAVEL</h1>
        <p>上传总名单 · 选日期 · 一键生成 Rooming List (与模板完全一致)</p>
    </div>
    """,
    unsafe_allow_html=True,
)


def sec(n, title, sub):
    st.markdown(
        f'<div class="sec"><div class="n">{n}</div>'
        f'<div class="t">{title}<small>{sub}</small></div></div>',
        unsafe_allow_html=True,
    )


# ============================================================
#  LOGICA EXCEL
# ============================================================
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "template_rooming.xlsx")

COLONNE_DATE = {
    "rooming": ("首晚入住日期", "末晚入住日期"),   # 排房表   → 2ª e 3ª colonna
    "passenger": ("上团日期", "下团日期"),          # 团员名单 → 1ª e 4ª colonna
}
COL_COGNOME = "姓"
COLONNE_TPL = ["姓", "名", "性别", "出生日期", "国籍", "护照号", "到期时间", "房型", "备注"]
COL_DATE_TPL = ("出生日期", "到期时间")
RIGA_DATI = 5
COL_A_K = range(1, 12)


def _pulisci_testo(v):
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    s = str(v)
    if s.strip().lower() in ("nan", "nat", "none"):
        return ""
    for ch in ["\xa0", " ", "　"]:
        s = s.replace(ch, " ")
    return s.strip()


def _fmt_data(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%d/%m/%Y")
    s = _pulisci_testo(v)
    if not s:
        return ""
    d = pd.to_datetime(s, errors="coerce", dayfirst=True)
    return d.strftime("%d/%m/%Y") if not pd.isna(d) else s


def leggi_master(file_bytes):
    """Legge il master escludendo le righe con celle SBARRATE (strikethrough)."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    intestazioni = [c.value for c in ws[1]]
    try:
        idx_cognome = intestazioni.index(COL_COGNOME)
    except ValueError:
        idx_cognome = None

    dati, n_barrate = [], 0
    for riga in ws.iter_rows(min_row=2):
        if idx_cognome is not None and riga[idx_cognome].value in (None, ""):
            continue
        barrata = any(
            cell.value not in (None, "") and cell.font and cell.font.strike
            for cell in riga
        )
        if barrata:
            n_barrate += 1
            continue
        dati.append([c.value for c in riga])
    return pd.DataFrame(dati, columns=intestazioni), n_barrate


def filtra_periodo(df, col_ini, col_fine, data_inizio, data_fine):
    ini = pd.to_datetime(df[col_ini], errors="coerce")
    fin = pd.to_datetime(df[col_fine], errors="coerce")
    mask = (ini <= pd.Timestamp(data_fine)) & (fin >= pd.Timestamp(data_inizio))
    return df[mask].copy()


def _snap(cell):
    return (copy.copy(cell.font), copy.copy(cell.fill), copy.copy(cell.border),
            copy.copy(cell.alignment), cell.number_format)


def _applica(cell, stile):
    f, fi, b, al, nf = stile
    cell.font = copy.copy(f)
    cell.fill = copy.copy(fi)
    cell.border = copy.copy(b)
    cell.alignment = copy.copy(al)
    cell.number_format = nf


def genera_da_template(tpl_bytes, df, hotel="", arriving_date=None):
    wb = load_workbook(io.BytesIO(tpl_bytes))
    ws = wb.active

    proto_dati = {c: _snap(ws.cell(row=RIGA_DATI, column=c)) for c in COL_A_K}
    proto_guida = {c: _snap(ws.cell(row=23, column=c)) for c in COL_A_K}
    proto_vuota = {c: _snap(ws.cell(row=30, column=c)) for c in COL_A_K}

    for rng in list(ws.merged_cells.ranges):
        s = str(rng)
        if s.startswith("I") and ":" in s:
            ws.unmerge_cells(s)

    for r in range(RIGA_DATI, 41):
        for c in COL_A_K:
            ws.cell(row=r, column=c).value = None

    ws["A3"].value = hotel or ""
    if arriving_date is not None:
        ws["F3"].value = f"Arriving Date {pd.Timestamp(arriving_date).strftime('%d/%m/%Y')}"

    n = len(df)
    r = RIGA_DATI
    for i, (_, riga) in enumerate(df.iterrows()):
        for c in COL_A_K:
            _applica(ws.cell(row=r, column=c), proto_dati[c])
        ws.cell(row=r, column=1, value=i + 1)
        for j, col in enumerate(COLONNE_TPL):
            v = riga.get(col)
            v = _fmt_data(v) if col in COL_DATE_TPL else _pulisci_testo(v)
            ws.cell(row=r, column=2 + j, value=v)
        ws.cell(row=r, column=11, value=None)
        r += 1

    for etichetta in ["GUIDA", "AUTISTA"]:
        for c in COL_A_K:
            _applica(ws.cell(row=r, column=c), proto_guida[c])
        ws.cell(row=r, column=1, value=r - RIGA_DATI + 1)
        ws.cell(row=r, column=2, value=etichetta)
        ws.cell(row=r, column=9, value="SGL")
        r += 1

    while r <= 34:
        for c in COL_A_K:
            _applica(ws.cell(row=r, column=c), proto_vuota[c])
        ws.cell(row=r, column=1, value=r - RIGA_DATI + 1)
        r += 1

    ws["M12"].value = n + 2  # Totale Pax = passeggeri + Guida + Autista
    for cell in ("M5", "M6", "M7", "M8", "M9"):
        ws[cell].value = None
    ws["M11"].value = "=SUM(M5:M9)"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------
#  ① Selettore modalità a carte
# ------------------------------------------------------------
if "mode" not in st.session_state:
    st.session_state.mode = "rooming"


def _pick(m):
    st.session_state.mode = m


sec("1", "选择报表类型", "Che tipo di elenco?")
c1, c2 = st.columns(2, gap="small")
c1.button(
    "🛏️\n排房表\nRooming list",
    key="m1",
    use_container_width=True,
    type="primary" if st.session_state.mode == "rooming" else "secondary",
    on_click=_pick,
    args=("rooming",),
)
c2.button(
    "👥\n团员名单\nPassenger list",
    key="m2",
    use_container_width=True,
    type="primary" if st.session_state.mode == "passenger" else "secondary",
    on_click=_pick,
    args=("passenger",),
)
modalita = st.session_state.mode
col_inizio, col_fine = COLONNE_DATE[modalita]
base_label = "首晚 / 末晚入住日期" if modalita == "rooming" else "上团 / 下团日期"

# ------------------------------------------------------------
#  ② Upload
# ------------------------------------------------------------
sec("2", "上传总名单", "Carica il master 报名表 (.xlsx)")
file_caricato = st.file_uploader("上传", type=["xlsx"], label_visibility="collapsed")


if file_caricato is not None:
    try:
        df_origine, n_barrate = leggi_master(file_caricato.getvalue())
        if n_barrate > 0:
            st.markdown(
                f'<div class="stat"><div class="big">🚫 {n_barrate}</div>'
                f'<div class="lbl">已自动忽略“划线删除(取消)”的乘客<br>'
                f'<b>有效乘客 {len(df_origine)}</b></div></div>',
                unsafe_allow_html=True,
            )

        if col_inizio not in df_origine.columns or col_fine not in df_origine.columns:
            st.error(
                f"❌ 此模式需要 '{col_inizio}' 和 '{col_fine}' 两列，未找到。"
                "请检查表头或切换上方模式。"
            )
            st.stop()

        serie_ini = pd.to_datetime(df_origine[col_inizio], errors="coerce")
        serie_fine = pd.to_datetime(df_origine[col_fine], errors="coerce")
        if not serie_ini.notna().any():
            st.error("⚠️ 日期列全部为空、错位或含乱码，无法筛选！")
            st.stop()

        min_date, max_date = serie_ini.min().date(), serie_fine.max().date()

        sec("3", "选择日期范围", f"基于 {base_label}")
        intervallo = st.date_input(
            "日期范围",
            value=(min_date, max_date),
            min_value=min_date,
            max_value=max_date,
            label_visibility="collapsed",
        )
        if not (isinstance(intervallo, tuple) and len(intervallo) == 2):
            st.warning("请在日历中选择完整的结束日期。")
            st.stop()

        d1, d2 = intervallo
        df_filtrato = filtra_periodo(df_origine, col_inizio, col_fine, d1, d2)
        st.markdown(
            f'<div class="stat"><div class="big">{len(df_filtrato)}</div>'
            f'<div class="lbl">位乘客符合条件<br><b>{d1}</b> → <b>{d2}</b></div></div>',
            unsafe_allow_html=True,
        )
        if df_filtrato.empty:
            st.info("当前没有符合条件的乘客。")
            st.stop()

        # --- ④ Info output + anteprima + download ---
        sec("4", "预览与下载", "Controlla e scarica")
        colonne_preview = [c for c in COLONNE_TPL if c in df_filtrato.columns]
        st.dataframe(df_filtrato[colonne_preview], use_container_width=True, hide_index=True)

        cha, chb = st.columns(2)
        hotel = cha.text_input("🏨 酒店名称 (Hotel)", value="")
        arriving = chb.date_input("📅 Arriving Date", value=d1)

        if modalita == "rooming":
            if not os.path.exists(TEMPLATE_PATH):
                st.error("缺少模板文件 template_rooming.xlsx，请将其与 streamlit_app.py 放在同一目录。")
                st.stop()
            with open(TEMPLATE_PATH, "rb") as f:
                tpl_bytes = f.read()
            output_bytes = genera_da_template(tpl_bytes, df_filtrato, hotel=hotel, arriving_date=arriving)
            nome_file = f"Rooming_List_{d1}_al_{d2}.xlsx"
            label = "📥 下载 Rooming List (与模板完全一致)"
        else:
            # modalità passenger: esporta i dati filtrati e puliti (formato semplice)
            df_exp = df_filtrato[colonne_preview].copy()
            for c in COL_DATE_TPL:
                if c in df_exp.columns:
                    df_exp[c] = df_exp[c].map(_fmt_data)
            buffer = io.BytesIO()
            with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                df_exp.to_excel(writer, index=False)
            output_bytes = buffer.getvalue()
            nome_file = f"团队行程名单_{d1}_al_{d2}.xlsx"
            label = "📥 下载 团员名单"

        st.markdown("<div style='height:6px'></div>", unsafe_allow_html=True)
        st.download_button(
            label=label,
            data=output_bytes,
            file_name=nome_file,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    except Exception as e:
        st.error(f"运行过程中发生错误：{e}")
else:
    st.markdown(
        '<div style="text-align:center;color:#a5aec2;padding:22px 0 4px;font-size:.92rem;">'
        "👆 上传总名单即可开始 · Carica il master per iniziare</div>",
        unsafe_allow_html=True,
    )
